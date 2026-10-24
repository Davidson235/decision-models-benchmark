"""Render a run summary into markdown and xlsx."""
from __future__ import annotations

import time
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _fmt(value, digits=3):
    if value is None:
        return "—"
    if isinstance(value, float):
        if value != value:  # NaN
            return "—"
        return f"{value:.{digits}f}"
    return str(value)


def _bar(value: float, width: int = 10) -> str:
    if value is None or value != value:
        return "—"
    filled = int(round(max(0.0, min(1.0, value)) * width))
    return "█" * filled + "·" * (width - filled)


def write_markdown(summary: dict, run: Path) -> Path:
    models = summary.get("models", {})
    usable = {m: e for m, e in models.items() if "accuracy" in e}
    ranked = sorted(usable.items(), key=lambda kv: kv[1]["accuracy"], reverse=True)
    lines: list[str] = []
    add = lines.append

    add(f"# Decision Models Benchmark — {summary.get('run', run.name)}")
    add("")
    add(f"Generated: `{summary.get('generated_at', time.strftime('%Y-%m-%dT%H:%M:%S'))}`  ")
    add(f"Paired tasks used for pairwise tests: **{summary.get('paired_tasks', 0)}**  ")
    add("Metric definitions: `p_true` is the probability assigned to the gold label; "
        "ECE uses 10 equal-width bins over the confidence of the chosen answer.")
    add("")

    add("## 1. Overall")
    add("")
    add("| # | Model | n | Accuracy | 95% CI | Brier ↓ | NLL ↓ | ECE ↓ | p50 ms | p95 ms | Cost $ | Out tok |")
    add("|---|-------|---|----------|--------|---------|-------|-------|--------|--------|--------|---------|")
    for i, (model, e) in enumerate(ranked, 1):
        lo, hi = e["accuracy_ci"]
        add(
            f"| {i} | `{model}` | {e['n']} | **{_fmt(e['accuracy'])}** | "
            f"{_fmt(lo,2)}–{_fmt(hi,2)} | {_fmt(e['brier'])} | {_fmt(e['nll'])} | "
            f"{_fmt(e['ece'])} | {_fmt(e['latency']['p50'],0)} | {_fmt(e['latency']['p95'],0)} | "
            f"{_fmt(e['cost_usd'],4)} | {e.get('output_tokens', 0)} |"
        )
    add("")
    add("> Accuracy differences below ~10 pp are not resolvable at n≈107; the CI column is the "
        "honest precision of each estimate, and section 5 gives paired tests.")
    add("")

    add("## 2. Accuracy by task family")
    add("")
    families = sorted({f for e in usable.values() for f in e["by_family"]})
    add("| Family | " + " | ".join(f"`{m.split('/')[-1]}`" for m, _ in ranked) + " |")
    add("|--------|" + "|".join(["------"] * len(ranked)) + "|")
    for fam in families:
        cells = []
        for _, e in ranked:
            v = e["by_family"].get(fam)
            cells.append(f"{_fmt(v['accuracy'])} ({v['n']})" if v else "—")
        add(f"| {fam} | " + " | ".join(cells) + " |")
    add("")

    add("## 3. Calibration")
    add("")
    for model, e in ranked:
        cal = e.get("calibration") or []
        if not cal:
            continue
        add(f"**{model}** — ECE {_fmt(e['ece'])}")
        add("")
        add("| bin | n | avg confidence | accuracy | gap |")
        add("|-----|---|----------------|----------|-----|")
        for row in cal:
            gap = row["accuracy"] - row["avg_conf"]
            add(f"| {row['bin']} | {row['n']} | {_fmt(row['avg_conf'])} | "
                f"{_fmt(row['accuracy'])} | {gap:+.3f} |")
        add("")

    add("## 4. Behavioural diagnostics")
    add("")
    add("| Model | Abstain on unanswerable | False abstain | RU/EN parity | Score MAE | nMAE | Bias |")
    add("|-------|------------------------|---------------|--------------|-----------|------|------|")
    for model, e in ranked:
        bias = e.get("score_bias")
        bias_cell = f"{bias:+.2f}" if bias is not None else "—"
        add(
            f"| `{model}` | {_fmt(e.get('abstention_rate_on_unanswerable'))} | "
            f"{_fmt(e.get('false_abstention_rate'))} | {_fmt(e.get('language_parity'))} "
            f"({e.get('twin_pairs', 0)} pairs) | {_fmt(e.get('score_mae'))} | "
            f"{_fmt(e.get('score_nmae'))} | {bias_cell} |"
        )
    add("")
    add("`Bias` is the mean of `predicted level - gold level` on the 20 score tasks. "
        "Every model lands below the rubric by 0.1–0.3 of a level. Two readings fit: the "
        "rubric anchors sit one notch higher than the written states justify, or models "
        "compress toward the middle of the scale. Note that no gold level in this corpus is "
        "4, so an empty level-4 bin is expected and does not by itself prove central "
        "tendency. Treat the score axis as agreement with this rubric, not as an absolute "
        "measurement of severity.")
    add("")
    add("| Model | predicted level → count |")
    add("|-------|------------------------|")
    for model, e in ranked:
        hist = e.get("score_pred_hist") or {}
        cells = ", ".join(f"{lvl}:{n}" for lvl, n in hist.items())
        add(f"| `{model}` | {cells or '—'} |")
    add("")
    add("")

    comps = summary.get("comparisons") or {}
    if comps:
        add("## 5. Paired comparison (bootstrap, 2000 resamples)")
        add("")
        add("| Pair (B − A) | Δ accuracy | 95% CI | p |")
        add("|--------------|-----------|--------|---|")
        for pair, res in sorted(comps.items(), key=lambda kv: -abs(kv[1]["diff"])):
            add(f"| {pair} | {res['diff']:+.3f} | {res['lo']:+.3f} … {res['hi']:+.3f} | {res['p']:.3f} |")
        add("")

    add("## 6. Reading of the numbers")
    add("")
    best = ranked[0] if ranked else None
    if best:
        bmodel, bentry = best
        add(f"- Highest raw accuracy: `{bmodel}` at {_fmt(bentry['accuracy'])} "
            f"(95% CI {_fmt(bentry['accuracy_ci'][0],2)}–{_fmt(bentry['accuracy_ci'][1],2)}).")
        cal_best = min(ranked, key=lambda kv: kv[1]["ece"])
        add(f"- Best calibrated: `{cal_best[0]}` with ECE {_fmt(cal_best[1]['ece'])}.")
        cheap = min(ranked, key=lambda kv: (kv[1]["cost_usd"], -kv[1]["accuracy"]))
        add(f"- Cheapest: `{cheap[0]}` at ${_fmt(cheap[1]['cost_usd'],4)} for the whole suite.")
        fast = min(ranked, key=lambda kv: kv[1]["latency"]["p50"] or 10**9)
        add(f"- Fastest p50: `{fast[0]}` at {_fmt(fast[1]['latency']['p50'],0)} ms.")
        add("")
        add("These four optima do not have to coincide; a model that wins on accuracy may lose "
            "on calibration, and the cheapest model may be the best calibrated one. The report "
            "keeps the axes separate on purpose — collapsing them into one number is a value "
            "judgement, not a measurement.")
    add("")

    path = run / "report.md"
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def _sheet(wb: Workbook, title: str, header: list[str], rows: list[list]):
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        ws.append(row)
    for i, name in enumerate(header, 1):
        width = max(len(str(name)) + 2, *(len(str(r[i - 1])) for r in rows)) if rows else len(name) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(44, max(10, width))
    ws.freeze_panes = "A2"
    return ws


def write_xlsx(summary: dict, run: Path) -> Path:
    models = summary.get("models", {})
    usable = {m: e for m, e in models.items() if "accuracy" in e}
    ranked = sorted(usable.items(), key=lambda kv: kv[1]["accuracy"], reverse=True)

    wb = Workbook()
    wb.remove(wb.active)

    rows = []
    for i, (model, e) in enumerate(ranked, 1):
        lo, hi = e["accuracy_ci"]
        rows.append([
            i, model, e["n"], e["accuracy"], lo, hi, e["brier"], e["nll"], e["ece"],
            e["latency"]["p50"], e["latency"]["p95"], e["cost_usd"],
            e.get("input_tokens", 0), e.get("output_tokens", 0),
            e.get("abstention_rate_on_unanswerable"), e.get("false_abstention_rate"),
            e.get("language_parity"), e.get("score_mae"), e.get("score_nmae"),
        ])
    _sheet(wb, "Overall", [
        "#", "model", "n", "accuracy", "ci_lo", "ci_hi", "brier", "nll", "ece",
        "p50_ms", "p95_ms", "cost_usd", "input_tokens", "output_tokens",
        "abstain_rate", "false_abstain", "ru_en_parity", "score_mae", "score_nmae",
    ], rows)

    fam_rows = []
    families = sorted({f for e in usable.values() for f in e["by_family"]})
    for fam in families:
        row = [fam]
        for _, e in ranked:
            v = e["by_family"].get(fam)
            row.append(v["accuracy"] if v else None)
        fam_rows.append(row)
    _sheet(wb, "ByFamily", ["family"] + [m for m, _ in ranked], fam_rows)

    cal_rows = []
    for model, e in ranked:
        for row in e.get("calibration") or []:
            cal_rows.append([model, row["bin"], row["n"], row["avg_conf"], row["accuracy"],
                             row["accuracy"] - row["avg_conf"]])
    _sheet(wb, "Calibration", ["model", "bin", "n", "avg_conf", "accuracy", "gap"], cal_rows)

    comp_rows = [[pair, r["diff"], r["lo"], r["hi"], r["p"]]
                 for pair, r in (summary.get("comparisons") or {}).items()]
    _sheet(wb, "Comparisons", ["pair", "delta_accuracy", "ci_lo", "ci_hi", "p"], comp_rows)

    path = run / "report.xlsx"
    wb.save(path)
    return path
